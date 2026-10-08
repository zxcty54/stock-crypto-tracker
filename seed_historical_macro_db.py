#!/usr/bin/env python3
"""Harvest official India monthly macro indicators for five infrastructure sectors.

The target window is the last 36 completed calendar months. This script does not
fill gaps, extrapolate, or relabel production/consumption as direct end-market
demand. A failed source is recorded and isolated; other sources continue.

Run:
    python -m pip install -r requirements_macro_seeder.txt
    python harvest_macro_historical_db.py
    python harvest_macro_historical_db.py --as-of 2026-10-08

The output defaults to ``macro_historical_db.json``. Source documents are retained
as official URLs plus SHA-256 hashes; downloaded files are not committed.
"""
from __future__ import annotations

import argparse
import copy
import hashlib
import json
import math
import re
import sys
import tempfile
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta, timezone
from html.parser import HTMLParser
from io import BytesIO
from pathlib import Path
from typing import Any, Callable, Iterable
from urllib.parse import urljoin, urlparse
from zoneinfo import ZoneInfo

import requests
from openpyxl import load_workbook
from pypdf import PdfReader

import seed_institutional_macro_db as oea

DEFAULT_OUTPUT = Path("macro_historical_db.json")
DEFAULT_MONTHS = 36
REQUEST_TIMEOUT_SECONDS = 60
USER_AGENT = "StockPulseInstitutionalMacroHarvester/2.0 (public Government of India data)"

PPAC_CONSUMPTION_PAGE = "https://ppac.gov.in/consumption/products-wise"
TRW_PUBLICATION_PAGE = "https://shipmin.gov.in/en/publication/trw-publication"
TRW_ARCHIVE_PAGE = "https://shipmin.gov.in/en/division/archieve/transport-research"

FISCAL_MONTHS = {
    "april": 4,
    "may": 5,
    "june": 6,
    "july": 7,
    "august": 8,
    "september": 9,
    "october": 10,
    "november": 11,
    "december": 12,
    "january": 1,
    "february": 2,
    "march": 3,
}
SHORT_MONTHS = {name[:3]: month for name, month in FISCAL_MONTHS.items()}

SECTOR_LABELS = {
    "cement": "Cement & Building Materials",
    "steel": "Steel",
    "road_epc": "Road EPC",
    "logistics": "Logistics",
    "port_exim": "Port / EXIM",
}

# Keep source meanings explicit. PPAC's HSD and bitumen series are proxies, not
# direct freight-demand or EPC-award measures. TRW overseas cargo is port cargo
# throughput (imports + exports), not a merchandise-trade value series.
SERIES_DEFINITIONS: dict[str, dict[str, Any]] = {
    "ici_cement_production_index": {
        "name": "Cement production volume index",
        "unit": "index (2022-23=100)",
        "measurement_type": "production_volume_index_not_end_use_demand",
        "source_id": "office_of_economic_adviser_ici",
        "definition": "Monthly production index; not cement consumption or end-user demand.",
    },
    "ici_steel_production_index": {
        "name": "Steel production volume index",
        "unit": "index (2022-23=100)",
        "measurement_type": "production_volume_index_not_end_use_demand",
        "source_id": "office_of_economic_adviser_ici",
        "definition": "Monthly production index; not finished-steel consumption or end-user demand.",
    },
    "ppac_bitumen_consumption": {
        "name": "Domestic bitumen consumption",
        "unit": "'000 metric tonnes",
        "measurement_type": "road_material_consumption_activity_proxy",
        "source_id": "ppac_petroleum_consumption",
        "definition": (
            "All-India domestic bitumen consumption. Road-material/activity proxy only; "
            "not Road EPC awards, execution, or contractor order inflows."
        ),
    },
    "ppac_hsd_consumption": {
        "name": "Domestic high-speed diesel (HSD) consumption",
        "unit": "'000 metric tonnes",
        "measurement_type": "fuel_consumption_activity_proxy_not_freight_demand",
        "source_id": "ppac_petroleum_consumption",
        "definition": (
            "All-India HSD consumption across end uses. Fuel/activity proxy only; "
            "not a direct freight-volume measure."
        ),
    },
    "mopsw_major_overseas_cargo": {
        "name": "Overseas cargo handled at Major Ports",
        "unit": "million tonnes",
        "measurement_type": "major_port_overseas_cargo_throughput",
        "source_id": "mopsw_trw_major_ports",
        "definition": (
            "Overseas cargo throughput reported by MoPSW for Major Ports; combines "
            "cargo loaded and unloaded and is not separated into exports and imports."
        ),
    },
    "mopsw_major_total_cargo": {
        "name": "Total cargo handled at Major Ports",
        "unit": "million tonnes",
        "measurement_type": "major_port_total_cargo_throughput_includes_coastal",
        "source_id": "mopsw_trw_major_ports",
        "definition": "Total Major Port cargo throughput, including overseas and coastal cargo.",
    },
    "mopsw_major_coastal_cargo": {
        "name": "Coastal cargo handled at Major Ports",
        "unit": "million tonnes",
        "measurement_type": "major_port_coastal_cargo_throughput",
        "source_id": "mopsw_trw_major_ports",
        "definition": "Coastal (domestic) cargo throughput at Major Ports.",
    },
}


class HarvesterError(RuntimeError):
    """Raised for a source-specific data or format problem."""


@dataclass
class SourceResult:
    id: str
    publisher: str
    series: str
    data_page_url: str
    definition: str
    status: str = "not_attempted"
    error: str | None = None
    errors: list[str] = field(default_factory=list)
    series_values: dict[str, dict[str, float | None]] = field(default_factory=dict)
    period_status: dict[str, dict[str, str]] = field(default_factory=dict)
    period_document_url: dict[str, dict[str, str]] = field(default_factory=dict)
    documents: list[dict[str, Any]] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)
    wpi_records: list[Any] = field(default_factory=list)

    def set_value(
        self,
        series_id: str,
        period: str,
        value: Any,
        *,
        status: str = "observed",
        document_url: str | None = None,
    ) -> None:
        self.series_values.setdefault(series_id, {})[period] = finite_number(value)
        self.period_status.setdefault(series_id, {})[period] = status
        if document_url:
            self.period_document_url.setdefault(series_id, {})[period] = document_url

    def add_document(
        self,
        url: str,
        title: str,
        file_format: str,
        *,
        periods: Iterable[str] = (),
        content: bytes | None = None,
    ) -> None:
        if not url:
            return
        existing = next((item for item in self.documents if item.get("url") == url), None)
        period_list = sorted(set(periods))
        digest = hashlib.sha256(content).hexdigest() if content is not None else None
        if existing:
            existing["periods"] = sorted(set(existing.get("periods", [])) | set(period_list))
            if digest and not existing.get("sha256"):
                existing["sha256"] = digest
            return
        self.documents.append(
            {
                "source_id": self.id,
                "title": title,
                "url": url,
                "format": file_format,
                "periods": period_list,
                "sha256": digest,
            }
        )


def finite_number(value: Any) -> float | None:
    number = oea._number(value)
    if number is None or not math.isfinite(number):
        return None
    return number


def previous_calendar_month(as_of: date) -> str:
    first_of_month = as_of.replace(day=1)
    last_complete = first_of_month - timedelta(days=1)
    return f"{last_complete.year:04d}-{last_complete.month:02d}"


def shift_period(period: str, month_delta: int) -> str:
    match = re.fullmatch(r"(\d{4})-(\d{2})", period)
    if not match:
        raise ValueError(f"Invalid YYYY-MM period: {period}")
    index = int(match.group(1)) * 12 + (int(match.group(2)) - 1) + month_delta
    year, month0 = divmod(index, 12)
    return f"{year:04d}-{month0 + 1:02d}"


def completed_month_window(as_of: date, months: int = DEFAULT_MONTHS) -> list[str]:
    """Return exactly the last N completed calendar months, ending last month."""
    if months < 1:
        raise ValueError("months must be a positive integer")
    end = previous_calendar_month(as_of)
    return [shift_period(end, offset) for offset in range(-(months - 1), 1)]


def india_today() -> date:
    return datetime.now(ZoneInfo("Asia/Kolkata")).date()


def _normalized_text(value: Any) -> str:
    text = unicodedata.normalize("NFKD", "" if value is None else str(value))
    text = "".join(char for char in text if not unicodedata.combining(char))
    text = text.replace("&", " and ")
    text = re.sub(r"[^a-zA-Z0-9]+", " ", text).strip().lower()
    return re.sub(r"\s+", " ", text)


def _month_number(value: Any) -> int | None:
    label = _normalized_text(value)
    if not label:
        return None
    token = label.split()[0]
    return FISCAL_MONTHS.get(token) or SHORT_MONTHS.get(token[:3])


def _parse_fiscal_year_start(text: str) -> int | None:
    patterns = (
        r"(?:FY|financial\s+year|fiscal\s+year)\s*[:'’\- ]*\s*(20\d{2})\s*[-/–]\s*(?:20)?\d{2}",
        r"period\s*:?\s*april\s+(20\d{2})\s*[-–/]\s*march\s+(?:20)?\d{2}",
        r"(20\d{2})\s*[-/–]\s*(?:20)?\d{2}",
    )
    for pattern in patterns:
        match = re.search(pattern, text, flags=re.IGNORECASE)
        if match:
            year = int(match.group(1))
            if 1990 <= year <= 2100:
                return year
    return None


def _fiscal_month_period(month: int, fiscal_start: int) -> str:
    year = fiscal_start if month >= 4 else fiscal_start + 1
    return f"{year:04d}-{month:02d}"


def _period_from_text(text: str) -> str | None:
    month_names = "January|February|March|April|May|June|July|August|September|October|November|December"
    month_abbreviations = "Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Sept|Oct|Nov|Dec"
    patterns = (
        rf"\b({month_names}|{month_abbreviations})\.?\s*,?\s*(20\d{{2}})\b",
        rf"\b({month_names}|{month_abbreviations})\.?\s*[-/]\s*(\d{{2}}|20\d{{2}})\b",
        r"\b(20\d{2})[-/](0[1-9]|1[0-2])\b",
    )
    for pattern in patterns:
        match = re.search(pattern, text, flags=re.IGNORECASE)
        if not match:
            continue
        if pattern.startswith(r"\b(20"):
            return f"{int(match.group(1)):04d}-{int(match.group(2)):02d}"
        month = _month_number(match.group(1))
        if month is None:
            continue
        year = int(match.group(2))
        if year < 100:
            year += 2000
        return f"{year:04d}-{month:02d}"
    return None


def _explicit_month_period(value: Any) -> str | None:
    period = oea.parse_month(value)
    if period:
        return period
    return _period_from_text(str(value)) if value is not None else None


def _status_for_window(source: SourceResult, target_periods: list[str]) -> None:
    count = sum(
        1
        for series_values in source.series_values.values()
        for period, value in series_values.items()
        if period in target_periods and value is not None
    )
    expected_series = {
        "office_of_economic_adviser_ici": [
            "ici_cement_production_index",
            "ici_steel_production_index",
        ],
        "office_of_economic_adviser_wpi": sorted(
            {
                request.series_id
                for sector_requests in oea.WPI_REQUESTS.values()
                for request in sector_requests
            }
        ),
        "ppac_petroleum_consumption": [
            "ppac_hsd_consumption",
            "ppac_bitumen_consumption",
        ],
        "mopsw_trw_major_ports": [
            "mopsw_major_total_cargo",
            "mopsw_major_overseas_cargo",
            "mopsw_major_coastal_cargo",
        ],
    }.get(source.id, list(source.series_values))
    if source.id == "office_of_economic_adviser_wpi":
        expected_series = sorted(set(expected_series) | {"wpi_high_speed_diesel"})
    target_slots = max(1, len(target_periods) * max(1, len(expected_series)))
    if count == 0:
        source.status = "failed" if source.error or source.errors else "unavailable"
    elif count >= target_slots:
        source.status = "available"
    else:
        source.status = "partial"


def new_source_results() -> dict[str, SourceResult]:
    return {
        "office_of_economic_adviser_ici": SourceResult(
            id="office_of_economic_adviser_ici",
            publisher="Office of the Economic Adviser, DPIIT, Government of India",
            series="Index of Core Industries (ICI), base year 2022-23",
            data_page_url=oea.ICI_INDEX_PAGE,
            definition="Monthly production volume indices; Cement and Steel are production, not end-use demand.",
        ),
        "office_of_economic_adviser_wpi": SourceResult(
            id="office_of_economic_adviser_wpi",
            publisher="Office of the Economic Adviser, DPIIT, Government of India",
            series="Wholesale Price Index (WPI), base year 2022-23",
            data_page_url=oea.WPI_INDEX_PAGE,
            definition="Wholesale price indicators; not company-specific invoices, tariffs, or contract prices.",
        ),
        "ppac_petroleum_consumption": SourceResult(
            id="ppac_petroleum_consumption",
            publisher="Petroleum Planning & Analysis Cell (PPAC), Ministry of Petroleum & Natural Gas, Government of India",
            series="Domestic consumption of petroleum products, product-wise",
            data_page_url=PPAC_CONSUMPTION_PAGE,
            definition=(
                "All-India market-demand estimates. Bitumen is a road-material/activity proxy; "
                "HSD is a fuel/activity proxy, not direct Road EPC or freight demand."
            ),
        ),
        "mopsw_trw_major_ports": SourceResult(
            id="mopsw_trw_major_ports",
            publisher="Transport Research Wing, Ministry of Ports, Shipping and Waterways, Government of India",
            series="Monthly cargo handling status for Major Ports",
            data_page_url=TRW_PUBLICATION_PAGE,
            definition=(
                "Major Port cargo throughput. Overseas cargo combines loaded and unloaded cargo; "
                "the series is not a separate export-value or import-value series."
            ),
        ),
    }


def _content_response(session: requests.Session, url: str) -> bytes:
    response = session.get(url, timeout=REQUEST_TIMEOUT_SECONDS)
    response.raise_for_status()
    content = response.content
    if not content:
        raise HarvesterError(f"Empty response from {url}")
    return content


def _new_session() -> requests.Session:
    session = requests.Session()
    session.headers.update(
        {
            "User-Agent": USER_AGENT,
            "Accept": "text/html,application/pdf,application/vnd.openxmlformats-officedocument.spreadsheetml.sheet,*/*",
        }
    )
    return session


def _collect_oea_ici(
    session: requests.Session,
    source: SourceResult,
    target_periods: list[str],
) -> None:
    url = oea.discover_latest_xlsx(session, oea.ICI_INDEX_PAGE, oea.ICI_FILE_PATTERN)
    content = oea.download_bytes(session, url)
    periods, values = oea.parse_ici_workbook(content)
    observed_periods = [period for period in periods if period in target_periods]
    source.add_document(
        url,
        "Index of Core Industries (ICI) monthly workbook, base year 2022-23",
        "xlsx",
        periods=periods,
        content=content,
    )
    for sector, series_id in (
        ("cement", "ici_cement_production_index"),
        ("steel", "ici_steel_production_index"),
    ):
        for period in periods:
            if period not in target_periods:
                continue
            raw_value = values.get(sector, {}).get(period)
            status = "provisional" if period == periods[-1] else "published_status_not_individually_verified"
            source.set_value(series_id, period, raw_value, status=status, document_url=url)
    if not observed_periods:
        source.error = "ICI workbook was found but had no monthly observations in the target window"
    source.notes.append(f"Workbook observations in target window: {len(observed_periods)}")


def _collect_oea_wpi(
    session: requests.Session,
    source: SourceResult,
    target_periods: list[str],
) -> None:
    url = oea.discover_latest_xlsx(session, oea.WPI_INDEX_PAGE, oea.WPI_FILE_PATTERN)
    content = oea.download_bytes(session, url)
    periods, records = oea.parse_wpi_workbook(content)
    source.wpi_records = records
    source.add_document(
        url,
        "Wholesale Price Index (WPI) monthly workbook, base year 2022-23",
        "xlsx",
        periods=periods,
        content=content,
    )
    requests_by_series: dict[str, oea.WPIRequest] = {}
    for sector_requests in oea.WPI_REQUESTS.values():
        for request in sector_requests:
            requests_by_series[request.series_id] = request
    # Logistics uses the same official HSD price row as the Road EPC input-cost context.
    hsd_request = next(
        request for request in oea.WPI_REQUESTS["road_epc"] if request.series_id == "wpi_high_speed_diesel"
    )
    requests_by_series[hsd_request.series_id] = hsd_request

    for series_id, request in requests_by_series.items():
        record = oea._choose_wpi_record(records, request)
        if record is None:
            source.notes.append(f"No exact WPI row found for {series_id}")
            continue
        for period in periods:
            if period not in target_periods:
                continue
            status = "provisional" if period in periods[-2:] else "published_status_not_individually_verified"
            source.set_value(series_id, period, record.values.get(period), status=status, document_url=url)
    observed_periods = [period for period in periods if period in target_periods]
    if not observed_periods:
        source.error = "WPI workbook was found but had no monthly observations in the target window"
    source.notes.append(f"Workbook observations in target window: {len(observed_periods)}")


@dataclass
class ParsedLink:
    href: str
    text: str
    attrs: dict[str, str]
    row_text: str = ""


class PPACPageParser(HTMLParser):
    """Read PPAC's server-rendered table and keep link/button metadata."""

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.text_parts: list[str] = []
        self.tables: list[list[list[str]]] = []
        self.links: list[ParsedLink] = []
        self.scripts: list[str] = []
        self._in_table = False
        self._table: list[list[str]] = []
        self._row: list[str] | None = None
        self._cell: list[str] | None = None
        self._link_attrs: dict[str, str] | None = None
        self._link_text: list[str] | None = None
        self._in_script = False
        self._script_text: list[str] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        attributes = {key.lower(): (value or "") for key, value in attrs}
        tag = tag.lower()
        if tag == "table":
            self._in_table = True
            self._table = []
        elif tag == "tr" and self._in_table:
            self._row = []
        elif tag in {"td", "th"} and self._row is not None:
            self._cell = []
        elif tag == "a" or tag == "button":
            self._link_attrs = attributes
            self._link_text = []
        if tag == "script":
            self._in_script = True
            self._script_text = []

    def handle_data(self, data: str) -> None:
        if data.strip():
            self.text_parts.append(data.strip())
        if self._cell is not None:
            self._cell.append(data)
        if self._link_text is not None:
            self._link_text.append(data)
        if self._in_script:
            self._script_text.append(data)

    def handle_endtag(self, tag: str) -> None:
        tag = tag.lower()
        if tag in {"td", "th"} and self._cell is not None and self._row is not None:
            self._row.append(" ".join(" ".join(self._cell).split()))
            self._cell = None
        elif tag == "tr" and self._row is not None and self._in_table:
            if self._row:
                self._table.append(self._row)
            self._row = None
        elif tag == "table" and self._in_table:
            if self._table:
                self.tables.append(self._table)
            self._table = []
            self._in_table = False
        if tag in {"a", "button"} and self._link_attrs is not None:
            self.links.append(
                ParsedLink(
                    href=self._link_attrs.get("href") or self._link_attrs.get("data-href") or "",
                    text=" ".join(" ".join(self._link_text or []).split()),
                    attrs=self._link_attrs,
                )
            )
            self._link_attrs = None
            self._link_text = None
        if tag == "script" and self._in_script:
            self.scripts.append("".join(self._script_text))
            self._script_text = []
            self._in_script = False


def _guess_page_fy_start(parser: PPACPageParser, as_of: date) -> int:
    text = " ".join(parser.text_parts)
    year = _parse_fiscal_year_start(text)
    if year is not None:
        return year
    # Fallback only determines which calendar year a month-only table refers to;
    # it does not create or estimate any measurement value.
    return as_of.year if as_of.month >= 4 else as_of.year - 1


def parse_ppac_html_table(
    html: str,
    *,
    as_of: date,
) -> tuple[dict[str, dict[str, float | None]], dict[str, str], list[ParsedLink], str]:
    parser = PPACPageParser()
    parser.feed(html)
    fiscal_start = _guess_page_fy_start(parser, as_of)
    values: dict[str, dict[str, float | None]] = {"ppac_hsd_consumption": {}, "ppac_bitumen_consumption": {}}
    period_status: dict[str, str] = {}
    for table in parser.tables:
        header_index: int | None = None
        month_columns: dict[int, tuple[int, str | None]] = {}
        for row_index, row in enumerate(table):
            candidate: dict[int, tuple[int, str | None]] = {}
            for column, cell in enumerate(row):
                explicit_period = _explicit_month_period(cell)
                month = int(explicit_period[-2:]) if explicit_period else _month_number(cell)
                if month is not None:
                    candidate[column] = (month, explicit_period)
            header_context = " ".join(
                _normalized_text(cell)
                for header_row in table[max(0, row_index - 2) : row_index + 1]
                for cell in header_row
            )
            has_product_row = any(
                (
                    _normalized_text(data_row[0]).startswith("hsd")
                    or "high speed diesel" in _normalized_text(data_row[0])
                    or _normalized_text(data_row[0]).startswith("bitumen")
                )
                for data_row in table[row_index + 1 :]
                if data_row
            )
            if len(candidate) >= 3 and ("product" in header_context or has_product_row):
                header_index = row_index
                month_columns = candidate
                break
        if header_index is None:
            continue
        for row in table[header_index + 1 :]:
            if not row:
                continue
            label = _normalized_text(row[0])
            if label == "hsd" or label.startswith("hsd ") or "high speed diesel" in label:
                series_id = "ppac_hsd_consumption"
            elif label == "bitumen" or label.startswith("bitumen "):
                series_id = "ppac_bitumen_consumption"
            else:
                continue
            for column, (month, explicit_period) in month_columns.items():
                if column >= len(row):
                    continue
                period = explicit_period or _fiscal_month_period(month, fiscal_start)
                values[series_id][period] = finite_number(row[column])
                period_status[period] = "provisional"
    visible_text = " ".join(parser.text_parts)
    return values, period_status, parser.links, visible_text


def _extract_download_candidates(html: str, links: list[ParsedLink]) -> tuple[list[str], list[str]]:
    historical: list[str] = []
    current: list[str] = []
    for link in links:
        attrs_text = " ".join(link.attrs.values())
        href = link.href
        payload = f"{link.text} {attrs_text} {href}"
        urls = [href] if href else []
        urls.extend(re.findall(r"(?:https?://[^\s\"'<>]+|/[^\s\"'<>]+)\.(?:xlsx|xls|pdf)(?:\?[^\s\"'<>]*)?", payload, flags=re.I))
        urls.extend(re.findall(r"(?:https?://[^\s\"'<>]*download\.php\?file=[^\s\"'<>]+|/download\.php\?file=[^\s\"'<>]+)", payload, flags=re.I))
        urls.extend(re.findall(r"[\"']((?:https?://|/)[^\s\"'<>]+)[\"']", attrs_text, flags=re.I))
        for raw_url in urls:
            if not raw_url or raw_url == "#" or raw_url.lower().startswith("javascript:"):
                continue
            normalized = raw_url.replace("&amp;", "&").rstrip(").,;")
            has_document_shape = bool(
                re.search(r"\.(?:xlsx|xls|pdf)(?:\?|$)|download\.php\?file=", normalized, re.I)
            )
            if "histor" in payload.lower():
                # PPAC's document route may be extensionless (for example, a
                # CMS /d/... URL); content type is checked after downloading.
                if has_document_shape or normalized.startswith(("/", "http://", "https://")):
                    historical.append(normalized)
            elif ("current" in payload.lower() or "pt consumption" in payload.lower()) and (
                has_document_shape or normalized.startswith(("/", "http://", "https://"))
            ):
                current.append(normalized)
    # PPAC pages sometimes keep download.php links in inline JavaScript rather
    # than an anchor. Preserve those URLs if nearby text says Historical.
    for match in re.finditer(r"(?:https?://[^\s\"'<>]*download\.php\?file=[^\s\"'<>]+|/download\.php\?file=[^\s\"'<>]+)", html, flags=re.I):
        raw_url = match.group(0).replace("&amp;", "&").rstrip(").,;")
        context = html[max(0, match.start() - 250) : min(len(html), match.end() + 250)].lower()
        if "histor" in context:
            historical.append(raw_url)
        elif "current" in context or "pt consumption" in context:
            current.append(raw_url)
    return list(dict.fromkeys(historical)), list(dict.fromkeys(current))


def _workbook_fiscal_start(workbook: Any, as_of: date) -> int:
    parts: list[str] = []
    for sheet in workbook.worksheets:
        parts.append(sheet.title)
        for row in sheet.iter_rows(values_only=True):
            parts.extend(str(value) for value in row if value is not None)
    return _parse_fiscal_year_start(" ".join(parts)) or (as_of.year if as_of.month >= 4 else as_of.year - 1)


def parse_ppac_workbook(
    content: bytes,
    *,
    as_of: date,
    target_periods: list[str],
) -> tuple[dict[str, dict[str, float | None]], dict[str, dict[str, str]]]:
    """Parse PPAC product/month workbook rows for HSD and Bitumen.

    Supports explicit month-year headers (Apr-23) and month-name headers within
    a fiscal-year sheet. Unparseable or blank cells remain null.
    """
    workbook = load_workbook(BytesIO(content), data_only=True, read_only=True)
    values: dict[str, dict[str, float | None]] = {"ppac_hsd_consumption": {}, "ppac_bitumen_consumption": {}}
    statuses: dict[str, dict[str, str]] = {"ppac_hsd_consumption": {}, "ppac_bitumen_consumption": {}}
    global_fy_start = _workbook_fiscal_start(workbook, as_of)
    try:
        for worksheet in workbook.worksheets:
            rows = list(worksheet.iter_rows(values_only=True))
            for header_index, row in enumerate(rows):
                month_columns: dict[int, tuple[int, str | None]] = {}
                for column, cell in enumerate(row):
                    explicit_period = _explicit_month_period(cell)
                    month = int(explicit_period[-2:]) if explicit_period else _month_number(cell)
                    if month is not None:
                        month_columns[column] = (month, explicit_period)
                if len(month_columns) < 3:
                    continue
                prior_text = " ".join(
                    str(value)
                    for prior_row in rows[max(0, header_index - 8) : header_index + 1]
                    for value in prior_row
                    if value is not None
                )
                fy_start = _parse_fiscal_year_start(worksheet.title + " " + prior_text) or global_fy_start
                for data_row in rows[header_index + 1 :]:
                    if not data_row:
                        continue
                    label = _normalized_text(data_row[0] if data_row else None)
                    if label == "hsd" or label.startswith("hsd ") or "high speed diesel" in label:
                        series_id = "ppac_hsd_consumption"
                    elif label == "bitumen" or label.startswith("bitumen "):
                        series_id = "ppac_bitumen_consumption"
                    else:
                        continue
                    for column, (month, explicit_period) in month_columns.items():
                        if column >= len(data_row):
                            continue
                        period = explicit_period or _fiscal_month_period(month, fy_start)
                        if period not in target_periods:
                            continue
                        parsed = finite_number(data_row[column])
                        # A later sheet/report version may revise an observation;
                        # preserve the last cell only when it is numeric.
                        if parsed is not None or period not in values[series_id]:
                            values[series_id][period] = parsed
                            statuses[series_id][period] = "provisional"
    finally:
        workbook.close()
    return values, statuses


def _pdf_text(content: bytes) -> str:
    if not content.startswith(b"%PDF"):
        raise HarvesterError("Expected a PDF document")
    reader = PdfReader(BytesIO(content), strict=False)
    pages = [page.extract_text() or "" for page in reader.pages]
    text = "\n".join(pages).strip()
    if not text:
        raise HarvesterError("PDF has no extractable text; OCR is not attempted")
    return text


def _parse_ppac_consumption_pdf(
    content: bytes,
    *,
    target_periods: list[str],
) -> tuple[str | None, dict[str, float | None]]:
    """Best-effort parser for a PPAC monthly product-consumption table."""
    text = _pdf_text(content).replace("\u00a0", " ")
    period = _period_from_text(text[:5000])
    if not period or period not in target_periods:
        return period, {}
    flat = re.sub(r"[ \t]+", " ", text)
    result: dict[str, float | None] = {}
    # Table 1 lists the preceding/current month before cumulative columns. Select
    # the current-month column (the second of the initial pair) only when the row
    # has the characteristic multi-column table, never from prose growth rates.
    for label, series_id in (("HSD", "ppac_hsd_consumption"), ("Bitumen", "ppac_bitumen_consumption")):
        candidates: list[float] = []
        for match in re.finditer(rf"\b{label}\b([^\n]{{0,220}})", flat, flags=re.IGNORECASE):
            tail = match.group(1)
            numbers = [finite_number(token) for token in re.findall(r"(?<![A-Za-z])\d[\d,]*(?:\.\d+)?", tail)]
            numbers = [number for number in numbers if number is not None]
            # Report tables use TMT, with HSD roughly 5,000-10,000 and Bitumen
            # roughly 100-1,500. Narrative MMT values are therefore rejected.
            plausible = [
                value
                for value in numbers
                if (1_000 <= value <= 15_000 if label == "HSD" else 50 <= value <= 2_000)
            ]
            if len(plausible) >= 2:
                candidates.append(plausible[1])
        if candidates:
            result[series_id] = candidates[0]
    return period, result


def _fetch_ppac_document(
    session: requests.Session,
    source: SourceResult,
    url: str,
    *,
    title: str,
    target_periods: list[str],
    as_of: date,
) -> None:
    absolute = urljoin(PPAC_CONSUMPTION_PAGE, url)
    if not _official_host(absolute):
        raise HarvesterError(f"Refusing non-Government PPAC document URL: {absolute}")
    content = _content_response(session, absolute)
    if content.startswith(b"PK\x03\x04"):
        values, statuses = parse_ppac_workbook(content, as_of=as_of, target_periods=target_periods)
        periods = sorted(
            set().union(*(set(series_values) for series_values in values.values()))
        )
        source.add_document(absolute, title, "xlsx", periods=periods, content=content)
        for series_id, series_values in values.items():
            for period, value in series_values.items():
                source.set_value(
                    series_id,
                    period,
                    value,
                    status=statuses.get(series_id, {}).get(period, "provisional"),
                    document_url=absolute,
                )
        return
    if content.startswith(b"%PDF"):
        period, values = _parse_ppac_consumption_pdf(content, target_periods=target_periods)
        if not period or not values:
            raise HarvesterError(f"Could not extract HSD/Bitumen values from {absolute}")
        source.add_document(absolute, title, "pdf", periods=[period], content=content)
        for series_id, value in values.items():
            source.set_value(series_id, period, value, status="provisional", document_url=absolute)
        return
    raise HarvesterError(f"PPAC download was neither XLSX nor PDF: {absolute}")


def harvest_ppac(
    session: requests.Session,
    source: SourceResult,
    target_periods: list[str],
    *,
    as_of: date,
) -> None:
    response = session.get(PPAC_CONSUMPTION_PAGE, timeout=REQUEST_TIMEOUT_SECONDS)
    response.raise_for_status()
    html = response.text
    page_values, page_status, links, _ = parse_ppac_html_table(html, as_of=as_of)
    for series_id, values in page_values.items():
        for period, value in values.items():
            if period in target_periods:
                source.set_value(series_id, period, value, status=page_status.get(period, "provisional"))

    historical_urls, current_urls = _extract_download_candidates(html, links)
    current_urls = [url for url in current_urls if "histor" not in url.lower()]
    historical_urls = [url for url in historical_urls if url not in current_urls]

    # Prefer one all-history document. PPAC's historical download may be gated by
    # its registration flow; a failed download is recorded, not treated as data.
    for url in historical_urls[:2]:
        try:
            _fetch_ppac_document(
                session,
                source,
                url,
                title="PPAC historical product-wise petroleum consumption report",
                target_periods=target_periods,
                as_of=as_of,
            )
        except Exception as exc:
            source.errors.append(f"Historical download {url}: {type(exc).__name__}: {exc}")

    # The current report is a backup for the current fiscal-year rows if the
    # server-rendered page contains values only for another view or lacks cells.
    for url in current_urls[:1]:
        try:
            _fetch_ppac_document(
                session,
                source,
                url,
                title="PPAC current product-wise petroleum consumption report",
                target_periods=target_periods,
                as_of=as_of,
            )
        except Exception as exc:
            source.errors.append(f"Current report {url}: {type(exc).__name__}: {exc}")

    for series_id, values in page_values.items():
        for period, value in values.items():
            if period in target_periods and value is not None:
                source.period_status.setdefault(series_id, {})[period] = page_status.get(period, "provisional")
    if not any(
        value is not None
        for values in source.series_values.values()
        for period, value in values.items()
        if period in target_periods
    ):
        source.error = "PPAC page was fetched, but no HSD/Bitumen observations could be read"
    if historical_urls:
        source.notes.append("Historical report link was found; download outcome is recorded above.")
    else:
        source.notes.append(
            "No public direct historical-file URL was exposed by the page parser. "
            "Rows not observed from the current report/page remain null."
        )
    if source.errors:
        source.error = source.error or "One or more PPAC document downloads failed"


@dataclass
class TRWDocument:
    period: str
    title: str
    url: str
    listing_url: str


class TRWListingParser(HTMLParser):
    """Collect TRW monthly publication links with their enclosing table rows."""

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.rows: list[tuple[str, list[tuple[str, str]]]] = []
        self.all_links: list[tuple[str, str]] = []
        self._in_row = False
        self._row_text: list[str] = []
        self._row_links: list[tuple[str, str]] = []
        self._in_cell = False
        self._cell_text: list[str] = []
        self._in_anchor = False
        self._anchor_href = ""
        self._anchor_text: list[str] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        tag = tag.lower()
        attributes = {key.lower(): value or "" for key, value in attrs}
        if tag == "tr":
            self._in_row = True
            self._row_text = []
            self._row_links = []
        elif tag in {"td", "th"} and self._in_row:
            self._in_cell = True
            self._cell_text = []
        elif tag == "a":
            self._in_anchor = True
            self._anchor_href = attributes.get("href", "")
            self._anchor_text = []

    def handle_data(self, data: str) -> None:
        if self._in_row:
            self._row_text.append(data)
        if self._in_cell:
            self._cell_text.append(data)
        if self._in_anchor:
            self._anchor_text.append(data)

    def handle_endtag(self, tag: str) -> None:
        tag = tag.lower()
        if tag in {"td", "th"} and self._in_cell:
            self._in_cell = False
        elif tag == "a" and self._in_anchor:
            link_text = " ".join(" ".join(self._anchor_text).split())
            if self._anchor_href:
                self.all_links.append((self._anchor_href, link_text))
                if self._in_row:
                    self._row_links.append((self._anchor_href, link_text))
            self._in_anchor = False
            self._anchor_href = ""
            self._anchor_text = []
        elif tag == "tr" and self._in_row:
            row_text = " ".join(" ".join(self._row_text).split())
            self.rows.append((row_text, self._row_links[:]))
            self._in_row = False
            self._row_text = []
            self._row_links = []


def _is_major_port_title(text: str) -> bool:
    normalized = _normalized_text(text)
    return (
        "major port" in normalized
        and "non major" not in normalized
        and "nonminor" not in normalized
        and "non minor" not in normalized
        and "hindi" not in normalized
    )


def parse_trw_listing(html: str, listing_url: str) -> list[TRWDocument]:
    parser = TRWListingParser()
    parser.feed(html)
    documents: list[TRWDocument] = []
    seen: set[str] = set()

    def add(period: str | None, title: str, href: str) -> None:
        if not period or not _is_major_port_title(title):
            return
        absolute_url = urljoin(listing_url, href)
        if not _official_host(absolute_url):
            return
        path = urlparse(absolute_url).path.lower()
        is_file = path.endswith((".pdf", ".doc", ".docx"))
        is_publication_detail = "/publication/" in path and absolute_url.rstrip("/") != listing_url.rstrip("/")
        if not (is_file or is_publication_detail) or absolute_url in seen:
            return
        seen.add(absolute_url)
        documents.append(TRWDocument(period, title, absolute_url, listing_url))

    # Table-based archive rows often hold a dated title and a separate Download
    # link; div/card views commonly make the title itself the anchor.
    for row_text, links in parser.rows:
        if not _is_major_port_title(row_text):
            continue
        period = _period_from_text(row_text)
        for href, link_text in links:
            add(period, row_text or link_text, href)

    for href, link_text in parser.all_links:
        absolute_url = urljoin(listing_url, href)
        path_title = urlparse(absolute_url).path.rsplit("/", 1)[-1].replace("-", " ").replace("_", " ")
        title = f"{link_text} {path_title}".strip()
        period = _period_from_text(title)
        add(period, title, href)
    return documents


def _extract_port_table_values(text: str, period: str) -> dict[str, float | None]:
    """Extract one target month's Major Port total/overseas/coastal tonnes."""
    month_number = int(period[-2:])
    month_name = next(name.title() for name, month in FISCAL_MONTHS.items() if month == month_number)
    year = int(period[:4])
    month_pattern = rf"{month_name}\s*,?\s*{year}"
    normalized = re.sub(r"\s+", " ", text.replace("\u00a0", " "))
    value_pattern = r"([+-]?(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d+)?)"
    unit_text = normalized.lower()
    if re.search(r"million\s+(?:metric\s+)?tonnes|\bmmt\b|\bmn\s+tonnes\b", unit_text):
        scale = 1.0
    elif re.search(r"thousand\s+(?:metric\s+)?tonnes|(?:'000|000)\s+(?:metric\s+)?tonnes|\bkt\b", unit_text):
        scale = 0.001
    elif re.search(r"\btonnes\b|\bmetric\s+tonnes\b", unit_text):
        scale = 0.000001
    else:
        scale = None
    values: dict[str, float | None] = {}
    occurrences = list(re.finditer(month_pattern, normalized, flags=re.IGNORECASE))
    for occurrence in occurrences:
        left = normalized[max(0, occurrence.start() - 500) : occurrence.start()]
        if "category" not in left.lower():
            continue
        snippet = normalized[occurrence.start() : occurrence.start() + 1700]
        labels = {
            "mopsw_major_total_cargo": r"\bTotal\b",
            "mopsw_major_overseas_cargo": r"\bOverseas\b",
            "mopsw_major_coastal_cargo": r"\bCoastal\b",
        }
        for series_id, label_pattern in labels.items():
            match = re.search(
                rf"{label_pattern}\s*[:\-]?\s*{value_pattern}(?:\s+{value_pattern})?",
                snippet,
                flags=re.IGNORECASE,
            )
            if not match:
                continue
            raw = finite_number(match.group(1))
            if raw is None:
                continue
            # Prefer the report's stated unit. If an extraction omits its table
            # header, retain plausible MMT values and normalize larger counts by
            # the only scale consistent with cargo throughput magnitudes.
            if scale is not None:
                mmt_value = raw * scale
            elif raw < 1_000:
                mmt_value = raw
            elif raw < 1_000_000:
                mmt_value = raw / 1_000
            else:
                mmt_value = raw / 1_000_000
            values[series_id] = mmt_value
        if values:
            return values

    # Fallback for line/column extraction where the table's month heading is
    # separated from row labels but a text summary states the three MMT figures.
    prose = normalized[occurrences[0].start() : occurrences[0].start() + 1400] if occurrences else normalized
    summary_patterns = {
        "mopsw_major_total_cargo": rf"Total Cargo handled\s*{value_pattern}\s*MMT",
        "mopsw_major_overseas_cargo": rf"Overseas(?: cargo)?\s*{value_pattern}\s*MMT",
        "mopsw_major_coastal_cargo": rf"Coastal(?: cargo)?\s*{value_pattern}\s*MMT",
    }
    for series_id, pattern in summary_patterns.items():
        match = re.search(pattern, prose, flags=re.IGNORECASE)
        if match:
            values[series_id] = finite_number(match.group(1))
    return values


def parse_trw_major_pdf(content: bytes, period: str) -> dict[str, float | None]:
    text = _pdf_text(content)
    return _extract_port_table_values(text, period)


def _resolve_trw_pdf_url(html_content: bytes, page_url: str) -> str:
    html = html_content.decode("utf-8", errors="replace")
    parser = TRWListingParser()
    parser.feed(html)
    links: list[tuple[str, str]] = []
    for href, title in parser.all_links:
        absolute = urljoin(page_url, href)
        if urlparse(absolute).path.lower().endswith(".pdf") and _official_host(absolute):
            links.append((absolute, title))
    if not links:
        for href in re.findall(r"href\s*=\s*[\"']([^\"']+\.pdf(?:\?[^\"']*)?)[\"']", html, flags=re.I):
            absolute = urljoin(page_url, href)
            if _official_host(absolute):
                links.append((absolute, ""))
    if not links:
        raise HarvesterError(f"No official PDF attachment was linked from {page_url}")
    links = list(dict.fromkeys(links))
    major_links = [
        item for item in links
        if _is_major_port_title(item[1] + " " + urlparse(item[0]).path)
    ]
    if len(major_links) == 1:
        return major_links[0][0]
    if len(links) == 1:
        return links[0][0]
    if major_links:
        return major_links[0][0]
    raise HarvesterError(f"Multiple PDFs found and no Major Ports attachment could be identified on {page_url}")


def _trw_listing_urls() -> list[str]:
    urls: list[str] = []
    # The newer publication view is paginated; the archive holds older months.
    for page in range(5):
        separator = "&" if "?" in TRW_PUBLICATION_PAGE else "?"
        urls.append(f"{TRW_PUBLICATION_PAGE}{separator}page={page}")
    urls.extend(
        [
            TRW_ARCHIVE_PAGE,
            "https://shipmin.gov.in/en/division/archieve/transport-research?page=1",
            "https://shipmin.gov.in/en/division/archieve/transport-research?page=2",
        ]
    )
    return urls


def harvest_trw_ports(
    session: requests.Session,
    source: SourceResult,
    target_periods: list[str],
) -> None:
    candidates: dict[str, list[TRWDocument]] = {period: [] for period in target_periods}
    successful_pages = 0
    for listing_url in _trw_listing_urls():
        try:
            response = session.get(listing_url, timeout=REQUEST_TIMEOUT_SECONDS)
            response.raise_for_status()
            docs = parse_trw_listing(response.text, listing_url)
            successful_pages += 1
            for doc in docs:
                if doc.period in candidates:
                    candidates[doc.period].append(doc)
        except Exception as exc:
            source.errors.append(f"Listing {listing_url}: {type(exc).__name__}: {exc}")

    # Prefer a document whose title explicitly says Major Ports; deduplicate the
    # current/archive copy by URL and download each month independently.
    for period in target_periods:
        docs = candidates.get(period, [])
        unique_docs = list({doc.url: doc for doc in docs}.values())
        unique_docs.sort(key=lambda doc: ("major ports" not in _normalized_text(doc.title), doc.url))
        parsed = False
        for doc in unique_docs:
            try:
                content = _content_response(session, doc.url)
                document_url = doc.url
                if not content.lstrip().startswith(b"%PDF"):
                    document_url = _resolve_trw_pdf_url(content, doc.url)
                    content = _content_response(session, document_url)
                if not content.lstrip().startswith(b"%PDF"):
                    raise HarvesterError("TRW report link and its attachment did not return a PDF")
                values = parse_trw_major_pdf(content, period)
                if not values:
                    raise HarvesterError(f"No monthly total/overseas/coastal table found for {period}")
                source.add_document(document_url, doc.title, "pdf", periods=[period], content=content)
                for series_id, value in values.items():
                    source.set_value(
                        series_id,
                        period,
                        value,
                        status="advance_estimate_or_published_status_not_individually_verified",
                        document_url=document_url,
                    )
                parsed = True
                break
            except Exception as exc:
                source.errors.append(f"{period} {doc.url}: {type(exc).__name__}: {exc}")
        if not parsed and not unique_docs:
            source.notes.append(f"No official Major Port monthly report was found in the TRW listings for {period}.")

    successful_months = sum(
        1
        for period in target_periods
        if any(
            source.series_values.get(series_id, {}).get(period) is not None
            for series_id in (
                "mopsw_major_total_cargo",
                "mopsw_major_overseas_cargo",
                "mopsw_major_coastal_cargo",
            )
        )
    )
    if successful_months == 0:
        source.error = "No target-month Major Port values were parsed from TRW reports"
    elif successful_months < len(target_periods):
        source.notes.append(
            f"Parsed {successful_months}/{len(target_periods)} target months; unavailable months remain null."
        )
    if successful_pages == 0:
        source.error = source.error or "All TRW listing pages failed"


def collect_sources(
    session: requests.Session,
    target_periods: list[str],
    *,
    as_of: date,
) -> dict[str, SourceResult]:
    sources = new_source_results()
    # Each source is attempted independently. A malformed or offline port report
    # cannot prevent ICI, WPI, or PPAC observations from being harvested.
    collectors: list[tuple[str, Callable[..., None], tuple[Any, ...], dict[str, Any]]] = [
        ("office_of_economic_adviser_ici", _collect_oea_ici, (target_periods,), {}),
        ("office_of_economic_adviser_wpi", _collect_oea_wpi, (target_periods,), {}),
        ("ppac_petroleum_consumption", harvest_ppac, (target_periods,), {"as_of": as_of}),
        ("mopsw_trw_major_ports", harvest_trw_ports, (target_periods,), {}),
    ]
    for source_id, collector, positional, keyword in collectors:
        source = sources[source_id]
        try:
            collector(session, source, *positional, **keyword)
        except Exception as exc:
            source.error = f"{type(exc).__name__}: {exc}"
        _status_for_window(source, target_periods)
    return sources


def _series_payload(
    series_id: str,
    source: SourceResult,
    target_periods: list[str],
    *,
    source_series: dict[str, Any] | None = None,
    role: str | None = None,
) -> dict[str, Any]:
    definition = SERIES_DEFINITIONS[series_id]
    values = source.series_values.get(series_id, {})
    period_status = source.period_status.get(series_id, {})
    period_document_url = source.period_document_url.get(series_id, {})
    observations: list[dict[str, Any]] = []
    for period in target_periods:
        value = finite_number(values.get(period))
        if period in period_status:
            observation_status = period_status[period]
        elif source.status == "failed":
            observation_status = "source_failed"
        elif source.status == "unavailable":
            observation_status = "source_unavailable"
        else:
            observation_status = "no_official_observation"
        observations.append(
            {
                "period": period,
                "value": value,
                "status": observation_status,
                "mom_change_pct": None,
                "yoy_change_pct": None,
                "source_page_url": source.data_page_url,
                "source_document_url": period_document_url.get(period),
            }
        )
    _add_actual_changes(observations)
    payload: dict[str, Any] = {
        "id": series_id,
        "name": definition["name"],
        "status": _series_status(observations, source.status),
        "measurement_type": definition["measurement_type"],
        "definition": definition["definition"],
        "unit": definition["unit"],
        "source_id": source.id,
        "source_url": source.data_page_url,
        "source_series": source_series,
        "observations": observations,
    }
    if role:
        payload["role"] = role
    return payload


def _series_status(observations: list[dict[str, Any]], source_status: str) -> str:
    numeric = sum(row["value"] is not None for row in observations)
    if numeric == len(observations):
        return "available"
    if numeric:
        return "partial"
    if source_status == "failed":
        return "source_failed_no_observations"
    return "no_observations_in_window"


def _add_actual_changes(observations: list[dict[str, Any]]) -> None:
    values = {row["period"]: finite_number(row.get("value")) for row in observations}
    for row in observations:
        period = row["period"]
        current = values.get(period)
        if current is None:
            continue
        for key, delta in (("mom_change_pct", -1), ("yoy_change_pct", -12)):
            prior = values.get(shift_period(period, delta))
            if prior is None or prior == 0:
                row[key] = None
            else:
                row[key] = round((current / prior - 1.0) * 100.0, 6)


def _wpi_source_record(source: SourceResult, request: oea.WPIRequest) -> dict[str, Any] | None:
    # WPI rows and metadata are retained from the current source result where
    # present; values themselves remain in the observation array.
    records = source.wpi_records
    record = oea._choose_wpi_record(records, request) if records else None
    if record is None:
        return None
    return {
        "level": record.level or None,
        "commodity_name": record.name,
        "commodity_code": record.code,
        "commodity_weight": record.weight,
    }


def _cost_series_for_sector(
    sector_id: str,
    wpi_source: SourceResult,
    target_periods: list[str],
) -> list[dict[str, Any]]:
    requests: list[oea.WPIRequest]
    if sector_id == "logistics":
        requests = [
            request
            for request in oea.WPI_REQUESTS["road_epc"]
            if request.series_id == "wpi_high_speed_diesel"
        ]
    else:
        requests = list(oea.WPI_REQUESTS.get(sector_id, ()))
    result: list[dict[str, Any]] = []
    for request in requests:
        record = _wpi_source_record(wpi_source, request)
        observations = _source_series_observations(
            request.series_id,
            wpi_source,
            target_periods,
            unit="index (2022-23=100)",
        )
        _add_actual_changes(observations)
        if record is None:
            status = "source_failed" if wpi_source.status == "failed" else "series_not_found_in_official_workbook"
        else:
            status = _series_status(observations, wpi_source.status)
        result.append(
            {
                "id": request.series_id,
                "name": request.display_name,
                "status": status,
                "role": request.role,
                "unit": "index (2022-23=100)",
                "source_series": record,
                "source_id": wpi_source.id,
                "source_url": wpi_source.data_page_url,
                "observations": observations,
            }
        )
    return result


def _source_series_observations(
    series_id: str,
    source: SourceResult,
    target_periods: list[str],
    *,
    unit: str | None = None,
) -> list[dict[str, Any]]:
    values = source.series_values.get(series_id, {})
    statuses = source.period_status.get(series_id, {})
    document_urls = source.period_document_url.get(series_id, {})
    result = []
    for period in target_periods:
        value = finite_number(values.get(period))
        if period in statuses:
            status = statuses[period]
        elif source.status == "failed":
            status = "source_failed"
        elif source.status == "unavailable":
            status = "source_unavailable"
        else:
            status = "no_official_observation"
        row = {
            "period": period,
            "value": value,
            "status": status,
            "mom_change_pct": None,
            "yoy_change_pct": None,
            "source_page_url": source.data_page_url,
            "source_document_url": document_urls.get(period),
        }
        if unit:
            row["unit"] = unit
        result.append(row)
    return result


def _demand_gap(metric_note: str) -> dict[str, Any]:
    return {
        "status": "not_populated_no_direct_monthly_demand_series",
        "metric": None,
        "unit": None,
        "observations": [],
        "note": metric_note,
    }


def build_database(
    *,
    generated_at: str,
    as_of: date,
    target_periods: list[str],
    source_results: dict[str, SourceResult],
) -> dict[str, Any]:
    ici = source_results["office_of_economic_adviser_ici"]
    wpi = source_results["office_of_economic_adviser_wpi"]
    ppac = source_results["ppac_petroleum_consumption"]
    trw = source_results["mopsw_trw_major_ports"]

    activity_by_sector: dict[str, list[dict[str, Any]]] = {
        "cement": [
            _series_payload("ici_cement_production_index", ici, target_periods),
        ],
        "steel": [
            _series_payload("ici_steel_production_index", ici, target_periods),
        ],
        "road_epc": [
            _series_payload(
                "ppac_bitumen_consumption",
                ppac,
                target_periods,
                role="road_material/activity proxy; not direct EPC demand or awards",
            ),
        ],
        "logistics": [
            _series_payload(
                "ppac_hsd_consumption",
                ppac,
                target_periods,
                role="fuel/activity proxy; not direct freight demand",
            ),
        ],
        "port_exim": [
            _series_payload("mopsw_major_overseas_cargo", trw, target_periods),
            _series_payload("mopsw_major_total_cargo", trw, target_periods),
            _series_payload("mopsw_major_coastal_cargo", trw, target_periods),
        ],
    }
    for series_list in activity_by_sector.values():
        for series in series_list:
            # Add changes only after all source rows are assembled.
            _add_actual_changes(series["observations"])

    sectors: dict[str, Any] = {}
    for sector_id, label in SECTOR_LABELS.items():
        if sector_id == "cement":
            demand_note = (
                "OEA's ICI Cement series is a production-volume index, not cement consumption. "
                "No direct monthly end-use demand observation is substituted."
            )
            context_refs: list[str] = []
        elif sector_id == "steel":
            demand_note = (
                "OEA's ICI Steel series is a production-volume index, not finished-steel consumption. "
                "No direct monthly end-use demand observation is substituted."
            )
            context_refs = []
        elif sector_id == "road_epc":
            demand_note = (
                "No direct monthly Road EPC award/execution series is included. PPAC bitumen "
                "consumption is retained separately as a road-material/activity proxy."
            )
            context_refs = [
                "cement.ici_cement_production_index",
                "steel.ici_steel_production_index",
            ]
        elif sector_id == "logistics":
            demand_note = (
                "PPAC HSD is all-India product consumption across end uses, not a direct freight "
                "or tonne-kilometre series. No freight demand is inferred from it."
            )
            context_refs = []
        else:
            demand_note = (
                "MoPSW overseas cargo throughput combines loading and unloading at Major Ports. "
                "It is not a separate merchandise-export or import series, and does not cover all ports."
            )
            context_refs = []

        sectors[sector_id] = {
            "code": sector_id.upper(),
            "name": label,
            "monthly_demand": _demand_gap(demand_note),
            "activity_indicators": activity_by_sector[sector_id],
            "context_indicator_refs": context_refs,
            "input_cost_indices": _cost_series_for_sector(sector_id, wpi, target_periods),
        }

    source_payloads: list[dict[str, Any]] = []
    all_documents: list[dict[str, Any]] = []
    source_errors: list[dict[str, Any]] = []
    for source in source_results.values():
        source_payloads.append(
            {
                "id": source.id,
                "publisher": source.publisher,
                "series": source.series,
                "definition": source.definition,
                "data_page": source.data_page_url,
                "status": source.status,
                "error": source.error,
                "errors": source.errors,
                "notes": source.notes,
                "documents": [doc["url"] for doc in source.documents],
                "observed_periods": sorted(
                    {
                        period
                        for values in source.series_values.values()
                        for period, value in values.items()
                        if value is not None and period in target_periods
                    }
                ),
            }
        )
        all_documents.extend(source.documents)
        if source.error or source.errors:
            source_errors.append(
                {
                    "source_id": source.id,
                    "error": source.error,
                    "details": source.errors,
                }
            )

    def observed_range(source: SourceResult) -> tuple[str | None, str | None]:
        periods = sorted(
            {
                period
                for series_values in source.series_values.values()
                for period, value in series_values.items()
                if value is not None and period in target_periods
            }
        )
        return (periods[0], periods[-1]) if periods else (None, None)

    data_period: dict[str, Any] = {
        "target_from": target_periods[0],
        "target_through": target_periods[-1],
        "months": len(target_periods),
        "by_source": {},
    }
    for source in source_results.values():
        start, end = observed_range(source)
        data_period["by_source"][source.id] = {"from": start, "through": end, "status": source.status}

    return {
        "schema_version": 2,
        "generated_at_utc": generated_at,
        "as_of_date_india": as_of.isoformat(),
        "country": "India",
        "frequency": "monthly",
        "scope": list(SECTOR_LABELS.keys()),
        "sector_codes": [value.upper() for value in SECTOR_LABELS],
        "window": {
            "months": len(target_periods),
            "from": target_periods[0],
            "through": target_periods[-1],
            "rule": "Last completed calendar month; current incomplete month excluded.",
        },
        "data_period": data_period,
        "sources": source_payloads,
        "documents": all_documents,
        "harvest_errors": source_errors,
        "sectors": sectors,
        "empirical_earnings_matrix": {
            "status": "not_computed",
            "observations": None,
            "note": "No earnings elasticity or company-level relationship is estimated by this harvester.",
        },
        "data_gaps": [
            {
                "id": "direct_demand_series",
                "status": "not_populated",
                "note": "Production and input/activity proxies are not relabelled as direct sector demand.",
            },
            {
                "id": "road_epc_awards_execution",
                "status": "no_direct_monthly_series_selected",
                "note": "Awards, value, constructed kilometres, and contractor order inflows are distinct metrics.",
            },
            {
                "id": "rail_freight_loading",
                "status": "not_included",
                "note": "No contiguous official national monthly rail series has been verified for this window; HSD remains only a fuel/activity proxy.",
            },
            {
                "id": "port_trade_direction",
                "status": "not_separated",
                "note": "TRW overseas cargo combines loadings and unloadings; import and export values are not separated.",
            },
        ],
    }


def _official_host(url: str | None) -> bool:
    if not url:
        return False
    host = (urlparse(url).hostname or "").lower()
    return host.endswith(".gov.in") or host.endswith(".nic.in") or host.endswith(".gov")


def _iter_series(sector: dict[str, Any]) -> Iterable[tuple[str, dict[str, Any]]]:
    for group in ("activity_indicators", "input_cost_indices"):
        for series in sector.get(group, []) or []:
            if isinstance(series, dict) and series.get("id"):
                yield group, series


def _trusted_prior_observation(observation: dict[str, Any], series_source_url: str | None) -> bool:
    status = _normalized_text(observation.get("status"))
    trusted_markers = ("observed", "actual", "published", "provisional", "official", "estimate")
    disallowed_markers = (
        "synthetic",
        "fabricated",
        "interpolated",
        "extrapolated",
        "seasonal",
        "cagr",
        "forecast",
        "projected",
        "fallback",
        "backfill",
        "modelled",
        "modeled",
    )
    if not any(marker in status for marker in trusted_markers):
        return False
    if any(marker in status for marker in disallowed_markers):
        return False
    source_urls = (
        observation.get("source_document_url"),
        observation.get("source_page_url"),
        observation.get("source_url"),
        series_source_url,
    )
    return any(_official_host(url) for url in source_urls)


def _trusted_prior_series(series: dict[str, Any]) -> bool:
    series_source_url = series.get("source_url") or series.get("data_page_url")
    numeric_observations = [
        row
        for row in series.get("observations", []) or []
        if isinstance(row, dict) and finite_number(row.get("value")) is not None
    ]
    return bool(numeric_observations) and all(
        _trusted_prior_observation(row, series_source_url) for row in numeric_observations
    )


def merge_previous_official_values(
    fresh: dict[str, Any],
    existing: dict[str, Any] | None,
    *,
    target_periods: list[str],
    source_results: dict[str, SourceResult],
) -> dict[str, Any]:
    """Preserve unknown DB fields and reuse only verified official cached rows.

    A prior official value is reused only for periods the current source did not
    observe, and only when the old observation has a Government of India source
    URL. Explicitly blank cells from a fetched document remain null.
    """
    if not isinstance(existing, dict):
        return fresh
    existing_sectors = existing.get("sectors")
    if not isinstance(existing_sectors, dict):
        existing_sectors = {}
    merged_sectors = copy.deepcopy(existing_sectors)
    for sector_id, new_sector in fresh["sectors"].items():
        old_sector = existing_sectors.get(sector_id, {})
        for group, new_series in _iter_series(new_sector):
            old_by_id = {
                old.get("id"): old
                for old_group, old in _iter_series(old_sector)
                if old_group == group
            }
            old_series = old_by_id.get(new_series.get("id"))
            if not old_series:
                continue
            source_id = new_series.get("source_id")
            source = source_results.get(source_id)
            if not source or source.status not in {"failed", "unavailable", "partial"}:
                continue
            old_source_url = old_series.get("source_url") or old_series.get("data_page_url")
            if not _official_host(old_source_url):
                continue
            old_observations = {
                row.get("period"): row
                for row in old_series.get("observations", [])
                if isinstance(row, dict) and row.get("period") in target_periods
            }
            for row in new_series.get("observations", []):
                if row.get("value") is not None:
                    continue
                if row.get("status") not in {
                    "source_failed",
                    "source_unavailable",
                    "no_official_observation",
                }:
                    continue
                previous = old_observations.get(row.get("period"))
                previous_value = finite_number(previous.get("value")) if previous else None
                if previous_value is None or previous is None:
                    continue
                if not _trusted_prior_observation(previous, old_source_url):
                    continue
                row["value"] = previous_value
                row["status"] = "cached_previous_official_observation"
                row["source_page_url"] = previous.get("source_page_url") or old_source_url
                row["source_document_url"] = previous.get("source_document_url")
                row["retained_from_existing_db"] = True
            _add_actual_changes(new_series.get("observations", []))
            new_series["status"] = _series_status(new_series.get("observations", []), source.status)
        # Keep old series outside the new known set, but do not replace fresh
        # observations for IDs that this harvester owns.
        for group in ("activity_indicators", "input_cost_indices"):
            old_list = old_sector.get(group, []) if isinstance(old_sector, dict) else []
            new_list = new_sector.get(group, [])
            new_ids = {item.get("id") for item in new_list if isinstance(item, dict)}
            for old in old_list or []:
                if (
                    isinstance(old, dict)
                    and old.get("id") not in new_ids
                    and _trusted_prior_series(old)
                ):
                    new_list.append(copy.deepcopy(old))
        merged_sectors[sector_id] = new_sector
    # Preserve any user-maintained sectors outside this harvester's fixed scope.
    fresh["sectors"] = merged_sectors

    # Preserve older source metadata that is not replaced by the current run.
    fresh_source_ids = {source.get("id") for source in fresh.get("sources", [])}
    old_sources = [
        source
        for source in existing.get("sources", []) or []
        if isinstance(source, dict) and source.get("id") not in fresh_source_ids
    ]
    fresh["sources"] = old_sources + fresh.get("sources", [])
    old_docs = [
        doc
        for doc in existing.get("documents", []) or []
        if isinstance(doc, dict) and doc.get("url") not in {item.get("url") for item in fresh.get("documents", [])}
    ]
    fresh["documents"] = old_docs + fresh.get("documents", [])

    document_by_url = {
        item.get("url"): item
        for item in fresh.get("documents", [])
        if isinstance(item, dict) and item.get("url")
    }
    current_source_by_id = {
        item.get("id"): item
        for item in fresh.get("sources", [])
        if isinstance(item, dict) and item.get("id")
    }
    old_source_by_id = {
        item.get("id"): item
        for item in existing.get("sources", []) or []
        if isinstance(item, dict) and item.get("id")
    }

    def retain_document(source_id: str, url: str, title: str, period: str | None = None) -> None:
        if not _official_host(url):
            return
        document = document_by_url.get(url)
        if document is None:
            suffix = Path(urlparse(url).path).suffix.lower().lstrip(".") or "unknown"
            document = {
                "source_id": source_id,
                "title": title,
                "url": url,
                "format": suffix,
                "periods": [],
                "sha256": None,
                "retained_from_existing_db": True,
            }
            fresh["documents"].append(document)
            document_by_url[url] = document
        if period:
            document["periods"] = sorted(set(document.get("periods", [])) | {period})
        source_payload = current_source_by_id.get(source_id)
        if source_payload is not None and url not in source_payload.get("documents", []):
            source_payload.setdefault("documents", []).append(url)

    for source_id, old_source in old_source_by_id.items():
        old_file_url = old_source.get("file_url")
        if old_file_url:
            retain_document(source_id, old_file_url, old_source.get("series", "Retained prior official source document"))
    for sector_id, sector in fresh.get("sectors", {}).items():
        for _, series in _iter_series(sector):
            for observation in series.get("observations", []) or []:
                if not observation.get("retained_from_existing_db"):
                    continue
                document_url = observation.get("source_document_url")
                if document_url:
                    retain_document(
                        series.get("source_id", "unknown"),
                        document_url,
                        series.get("name", "Retained prior official source document"),
                        observation.get("period"),
                    )
    # Keep arbitrary user-defined metadata not owned by this harvester.
    for key, value in existing.items():
        if key not in fresh and key not in {"sectors", "sources", "documents"}:
            fresh[key] = copy.deepcopy(value)
    return fresh


def _load_existing(path: Path) -> dict[str, Any] | None:
    if not path.exists():
        return None
    try:
        with path.open("r", encoding="utf-8") as handle:
            payload = json.load(handle)
    except (OSError, json.JSONDecodeError) as exc:
        raise HarvesterError(f"Existing database could not be read; leaving it untouched: {exc}") from exc
    if not isinstance(payload, dict):
        raise HarvesterError("Existing database root is not a JSON object; leaving it untouched")
    return payload


def _atomic_write_json(path: Path, payload: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(
        "w",
        encoding="utf-8",
        dir=path.parent,
        prefix=f".{path.name}.",
        suffix=".tmp",
        delete=False,
    ) as temporary:
        json.dump(payload, temporary, indent=2, ensure_ascii=False, allow_nan=False)
        temporary.write("\n")
        temporary.flush()
        temp_path = Path(temporary.name)
    temp_path.replace(path)


def seed(
    output_path: Path = DEFAULT_OUTPUT,
    *,
    as_of: date | None = None,
    session: requests.Session | None = None,
) -> dict[str, Any]:
    as_of = as_of or india_today()
    target_periods = completed_month_window(as_of, DEFAULT_MONTHS)
    owns_session = session is None
    active_session = session or _new_session()
    try:
        source_results = collect_sources(active_session, target_periods, as_of=as_of)
    finally:
        if owns_session:
            active_session.close()

    payload = build_database(
        generated_at=datetime.now(timezone.utc).isoformat(timespec="seconds"),
        as_of=as_of,
        target_periods=target_periods,
        source_results=source_results,
    )
    existing = _load_existing(output_path)
    payload = merge_previous_official_values(
        payload,
        existing,
        target_periods=target_periods,
        source_results=source_results,
    )
    _atomic_write_json(output_path, payload)
    counts = {
        source.id: sum(
            value is not None
            for series_values in source.series_values.values()
            for period, value in series_values.items()
            if period in target_periods
        )
        for source in source_results.values()
    }
    print(f"Wrote {output_path} for {target_periods[0]} through {target_periods[-1]} ({len(target_periods)} months).")
    for source_id, count in counts.items():
        print(f"  {source_id}: {source_results[source_id].status}, {count} numeric observations")
    return payload


def _parse_as_of(value: str | None) -> date | None:
    if not value:
        return None
    try:
        return date.fromisoformat(value)
    except ValueError as exc:
        raise argparse.ArgumentTypeError("--as-of must be YYYY-MM-DD") from exc


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--output",
        type=Path,
        default=DEFAULT_OUTPUT,
        help=f"JSON output path (default: {DEFAULT_OUTPUT})",
    )
    parser.add_argument(
        "--as-of",
        type=_parse_as_of,
        default=None,
        help="Override India-local current date for reproducible runs (YYYY-MM-DD)",
    )
    args = parser.parse_args(argv)
    try:
        seed(args.output, as_of=args.as_of)
    except (requests.RequestException, HarvesterError, OSError, ValueError) as exc:
        print(f"Harvester could not write output; existing database was left untouched: {exc}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
